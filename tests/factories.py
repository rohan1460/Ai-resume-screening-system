"""Builders for real PDF/DOCX bytes, so parser tests exercise real file formats."""

import io

import pymupdf
from docx import Document


def make_pdf_bytes(text: str) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    # insert_textbox wraps long lines; a bare insert_text would run off the page.
    rect = pymupdf.Rect(50, 50, 550, 780)
    page.insert_textbox(rect, text, fontsize=11)
    data: bytes = doc.tobytes()
    doc.close()
    return data


def make_two_column_pdf_bytes(
    left: list[str],
    right: list[str],
    header: str | None = None,
    footer: str | None = None,
) -> bytes:
    """A genuinely column-flowed PDF.

    Each line is drawn individually, which is how real two-column PDFs are laid out —
    and which makes PyMuPDF merge the two columns into one block per row.
    """
    doc = pymupdf.open()
    page = doc.new_page()
    y = 60.0
    if header:
        page.insert_text((50, 40), header, fontsize=14)
    for index in range(max(len(left), len(right))):
        row_y = y + index * 20
        if index < len(left):
            page.insert_text((50, row_y), left[index], fontsize=10)
        if index < len(right):
            page.insert_text((320, row_y), right[index], fontsize=10)
    if footer:
        page.insert_text((50, y + max(len(left), len(right)) * 20 + 30), footer, fontsize=12)
    data: bytes = doc.tobytes()
    doc.close()
    return data


def make_right_aligned_pdf_bytes(entries: list[tuple[str, str]], intro: list[str]) -> bytes:
    """Single-column text with a right-aligned date rail.

    This must NOT be mistaken for a two-column layout: the dates belong beside their
    entries, not collected together at the end.
    """
    doc = pymupdf.open()
    page = doc.new_page()
    for index, line in enumerate(intro):
        page.insert_text((50, 60 + index * 18), line, fontsize=10)
    start = 60 + len(intro) * 18 + 20
    for index, (label, date) in enumerate(entries):
        row_y = start + index * 22
        page.insert_text((50, row_y), label, fontsize=10)
        page.insert_text((430, row_y), date, fontsize=10)
    data: bytes = doc.tobytes()
    doc.close()
    return data


def make_table_pdf_bytes(rows: list[list[str]], title: str = "SKILLS MATRIX") -> bytes:
    """A PDF with a ruled table, so PyMuPDF's table detection has lines to find."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), title, fontsize=13)

    top, left = 80.0, 50.0
    row_h, col_w = 24.0, 160.0
    cols = len(rows[0])

    for r in range(len(rows) + 1):
        y = top + r * row_h
        page.draw_line(pymupdf.Point(left, y), pymupdf.Point(left + cols * col_w, y))
    for c in range(cols + 1):
        x = left + c * col_w
        page.draw_line(pymupdf.Point(x, top), pymupdf.Point(x, top + len(rows) * row_h))

    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            page.insert_text((left + c * col_w + 5, top + r * row_h + 16), value, fontsize=9)

    data: bytes = doc.tobytes()
    doc.close()
    return data


def make_docx_bytes(text: str, table_rows: list[list[str]] | None = None) -> bytes:
    document = Document()
    for line in text.split("\n"):
        document.add_paragraph(line)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.rows[r].cells[c].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


SAMPLE_RESUME = """Priya Sharma
Senior Backend Engineer
priya.sharma@example.com | +91 98765 43210

SUMMARY
Backend engineer with 6 years building distributed systems.

SKILLS
Python, FastAPI, PostgreSQL, Docker, Kubernetes, Redis, SQL

EXPERIENCE
Senior Backend Engineer, Acme Corp
Jan 2021 - Present
Built payment services handling 2M requests per day.

Backend Engineer, Globex
Jun 2018 - Dec 2020
Maintained internal APIs and data pipelines.

EDUCATION
B.Tech in Computer Science, IIT Bombay, 2018
"""
